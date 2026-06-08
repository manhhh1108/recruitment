"""
Task 7: Add Class Diagram (3.2.4) and Data Dictionary (3.2.5) sections
before "3.3. Thiet ke kiem soat".
"""

from copy import deepcopy
from docx import Document
from docx.oxml.ns import qn
from lxml import etree

DOC_PATH = "D:/JobPotal/recruitment_web-xw1rpa/recruitment_web/K62_211213818_BuiDucTaiAnh_CNTT4.docx"

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def find_template(doc, style_name):
    for p in doc.paragraphs:
        if p.style.name == style_name and p.runs:
            return p
    for p in doc.paragraphs:
        if p.style.name == style_name:
            return p
    return None


def make_styled_paragraph(doc, text, style_name, bold=False):
    """Create a paragraph with a given style by deep-copying a template."""
    template = find_template(doc, style_name)
    if template is not None:
        p = deepcopy(template._element)
        # Remove all children except pPr
        for child in list(p):
            tag = child.tag.split('}')[-1]
            if tag != 'pPr':
                p.remove(child)
        # Add run
        r = etree.SubElement(p, f'{{{W}}}r')
        if template.runs:
            rPr_src = template.runs[0]._element.find(f'{{{W}}}rPr')
            if rPr_src is not None:
                rPr = deepcopy(rPr_src)
                r.insert(0, rPr)
                if bold:
                    b_elem = rPr.find(f'{{{W}}}b')
                    if b_elem is None:
                        etree.SubElement(rPr, f'{{{W}}}b')
        elif bold:
            rPr = etree.SubElement(r, f'{{{W}}}rPr')
            r.insert(0, rPr)
            etree.SubElement(rPr, f'{{{W}}}b')
        t = etree.SubElement(r, f'{{{W}}}t')
        t.text = text
        t.set(f'{{{W}}}space', "preserve")
    else:
        p = etree.Element(f'{{{W}}}p')
        pPr = etree.SubElement(p, f'{{{W}}}pPr')
        pStyle = etree.SubElement(pPr, f'{{{W}}}pStyle')
        pStyle.set(f'{{{W}}}val', style_name)
        r = etree.SubElement(p, f'{{{W}}}r')
        if bold:
            rPr = etree.SubElement(r, f'{{{W}}}rPr')
            r.insert(0, rPr)
            etree.SubElement(rPr, f'{{{W}}}b')
        t = etree.SubElement(r, f'{{{W}}}t')
        t.text = text
        t.set(f'{{{W}}}space', "preserve")
    return p


def make_normal(doc, text):
    return make_styled_paragraph(doc, text, "Normal")


def make_heading3(doc, text):
    return make_styled_paragraph(doc, text, "Heading 3")


def make_hinh(doc, text):
    return make_styled_paragraph(doc, text, "hinh")


def make_bang(doc, text):
    return make_styled_paragraph(doc, text, u"B\u1ea3ng")


def make_blank(doc):
    template = find_template(doc, "Normal")
    if template:
        p = deepcopy(template._element)
        for child in list(p):
            tag = child.tag.split('}')[-1]
            if tag != 'pPr':
                p.remove(child)
        return p
    p = etree.Element(f'{{{W}}}p')
    return p


def make_table(headers, rows, col_widths=None):
    """Create a table XML element with Table Grid style and borders."""
    ncols = len(headers)

    tbl = etree.Element(f'{{{W}}}tbl')

    # tblPr
    tblPr = etree.SubElement(tbl, f'{{{W}}}tblPr')
    tblStyle = etree.SubElement(tblPr, f'{{{W}}}tblStyle')
    tblStyle.set(f'{{{W}}}val', "TableGrid")
    tblW = etree.SubElement(tblPr, f'{{{W}}}tblW')
    tblW.set(f'{{{W}}}w', "5000")
    tblW.set(f'{{{W}}}type', "pct")
    tblLook = etree.SubElement(tblPr, f'{{{W}}}tblLook')
    tblLook.set(f'{{{W}}}val', "04A0")
    tblLook.set(f'{{{W}}}firstRow', "1")
    tblLook.set(f'{{{W}}}lastRow', "0")
    tblLook.set(f'{{{W}}}firstColumn', "1")
    tblLook.set(f'{{{W}}}lastColumn', "0")
    tblLook.set(f'{{{W}}}noHBand', "0")
    tblLook.set(f'{{{W}}}noVBand', "1")

    # tblGrid
    tblGrid = etree.SubElement(tbl, f'{{{W}}}tblGrid')
    if col_widths is None:
        total = 9072  # ~16cm in twips
        w = total // ncols
        col_widths = [w] * ncols
    for cw in col_widths:
        gridCol = etree.SubElement(tblGrid, f'{{{W}}}gridCol')
        gridCol.set(f'{{{W}}}w', str(cw))

    def make_cell(text, is_bold=False):
        tc = etree.Element(f'{{{W}}}tc')
        p = etree.SubElement(tc, f'{{{W}}}p')
        r = etree.SubElement(p, f'{{{W}}}r')
        rPr = etree.SubElement(r, f'{{{W}}}rPr')
        r.insert(0, rPr)
        sz = etree.SubElement(rPr, f'{{{W}}}sz')
        sz.set(f'{{{W}}}val', "24")  # 12pt
        szCs = etree.SubElement(rPr, f'{{{W}}}szCs')
        szCs.set(f'{{{W}}}val', "24")
        if is_bold:
            etree.SubElement(rPr, f'{{{W}}}b')
        t = etree.SubElement(r, f'{{{W}}}t')
        t.text = text
        t.set(f'{{{W}}}space', "preserve")
        return tc

    # Header row
    tr_h = etree.SubElement(tbl, f'{{{W}}}tr')
    for h in headers:
        tr_h.append(make_cell(h, is_bold=True))

    # Data rows
    for row in rows:
        tr = etree.SubElement(tbl, f'{{{W}}}tr')
        for cell_text in row:
            tr.append(make_cell(cell_text, is_bold=False))

    return tbl


def main():
    doc = Document(DOC_PATH)

    # Find the anchor: "3.3. Thiet ke kiem soat" heading
    anchor = None
    for i, para in enumerate(doc.paragraphs):
        if para.style.name == 'Heading 2' and u'3.3. Thi\u1ebft k\u1ebf ki\u1ec3m so\u00e1t' in para.text:
            anchor = para._element
            print(f"Found anchor at paragraph {i}: {para.text}")
            break

    if anchor is None:
        raise RuntimeError("Could not find '3.3. Thiet ke kiem soat' heading")

    # Find the element just before anchor (the empty Normal paragraph)
    body = doc.element.body
    prev_elem = anchor.getprevious()
    print(f"Previous element tag: {prev_elem.tag.split('}')[-1] if prev_elem is not None else 'None'}")

    # We'll insert using addnext() on the RM figure paragraph (paragraph 536)
    # Actually, let's find "Hinh 3.2: So do RM" paragraph
    rm_para = None
    for para in doc.paragraphs:
        if u'S\u01a1 \u0111\u1ed3 RM' in para.text and para.style.name == 'hinh':
            rm_para = para._element
            print(f"Found RM figure paragraph: {para.text}")
            break

    if rm_para is None:
        raise RuntimeError("Could not find RM diagram figure paragraph")

    # Build all elements in forward order
    elements = []

    # === Section 3.2.4: Class Diagram ===
    elements.append(make_blank(doc))
    elements.append(make_heading3(doc, u"3.2.4. Bi\u1ec3u \u0111\u1ed3 l\u1edbp (Class Diagram)"))
    elements.append(make_normal(doc,
        u"Bi\u1ec3u \u0111\u1ed3 l\u1edbp m\u00f4 t\u1ea3 c\u1ea5u tr\u00fac c\u00e1c l\u1edbp Model trong h\u1ec7 th\u1ed1ng Backend Laravel v\u00e0 m\u1ed1i quan h\u1ec7 gi\u1eefa ch\u00fang. M\u1ed7i l\u1edbp Model t\u01b0\u01a1ng \u1ee9ng v\u1edbi m\u1ed9t b\u1ea3ng trong c\u01a1 s\u1edf d\u1eef li\u1ec7u v\u00e0 \u0111\u1ecbnh ngh\u0129a c\u00e1c thu\u1ed9c t\u00ednh (attributes) c\u00f9ng c\u00e1c m\u1ed1i quan h\u1ec7 (relationships) v\u1edbi c\u00e1c Model kh\u00e1c."))

    # Class diagram table
    class_headers = [u"T\u00ean l\u1edbp (Model)", u"B\u1ea3ng CSDL", u"Quan h\u1ec7 ch\u00ednh"]
    class_rows = [
        ["User", "users", "hasOne(Candidate), hasOne(Employer)"],
        ["Candidate", "candidates", "belongsTo(User), hasMany(Education, Experience, Skill, Project, Certificate, Prize, Activities, Other, Resume, CandidateMessage), belongsToMany(Job)"],
        ["Employer", "employers", "belongsTo(User), hasMany(Job), belongsToMany(Location)"],
        ["Job", "jobs", "belongsTo(Employer, Jtype, Jlevel), belongsToMany(Industry, Location, Jskill, Jtag, Candidate)"],
        ["Resume", "resumes", "belongsTo(Candidate)"],
        ["Education", "educations", "belongsTo(Candidate)"],
        ["Experience", "experiences", "belongsTo(Candidate), belongsTo(Jtype)"],
        ["Skill", "skills", "belongsTo(Candidate)"],
        ["Project", "projects", "belongsTo(Candidate)"],
        ["Certificate", "certificates", "belongsTo(Candidate)"],
        ["Prize", "prizes", "belongsTo(Candidate)"],
        ["Activities", "activities", "belongsTo(Candidate)"],
        ["Other", "others", "belongsTo(Candidate)"],
        ["Industry", "industries", "belongsToMany(Job)"],
        ["Location", "locations", "belongsToMany(Job), belongsToMany(Employer)"],
        ["Jtype", "jtypes", "hasMany(Job), hasMany(Experience)"],
        ["Jlevel", "jlevels", "hasMany(Job)"],
        ["Jskill", "jskills", "belongsToMany(Job)"],
        ["Jtag", "jtags", "belongsToMany(Job)"],
        ["CandidateMessage", "candidate_messages", "belongsTo(Candidate), belongsTo(Job)"],
    ]
    elements.append(make_table(class_headers, class_rows, col_widths=[2000, 2000, 5072]))
    elements.append(make_bang(doc, u"B\u1ea3ng 3.13: Danh s\u00e1ch c\u00e1c l\u1edbp Model v\u00e0 quan h\u1ec7"))
    elements.append(make_hinh(doc, u"(Ch\u00e8n bi\u1ec3u \u0111\u1ed3 l\u1edbp t\u1ea1i \u0111\u00e2y)"))
    elements.append(make_hinh(doc, u"H\u00ecnh 3.3: Bi\u1ec3u \u0111\u1ed3 l\u1edbp (Class Diagram) c\u1ee7a h\u1ec7 th\u1ed1ng"))

    # === Section 3.2.5: Data Dictionary ===
    elements.append(make_blank(doc))
    elements.append(make_heading3(doc, u"3.2.5. T\u1eeb \u0111i\u1ec3n d\u1eef li\u1ec7u (Data Dictionary)"))
    elements.append(make_normal(doc,
        u"T\u1eeb \u0111i\u1ec3n d\u1eef li\u1ec7u m\u00f4 t\u1ea3 chi ti\u1ebft c\u1ea5u tr\u00fac c\u00e1c b\u1ea3ng ch\u00ednh trong c\u01a1 s\u1edf d\u1eef li\u1ec7u h\u1ec7 th\u1ed1ng, bao g\u1ed3m t\u00ean c\u1ed9t, ki\u1ec3u d\u1eef li\u1ec7u, m\u00f4 t\u1ea3 v\u00e0 r\u00e0ng bu\u1ed9c."))

    dd_headers = [u"T\u00ean c\u1ed9t", u"Ki\u1ec3u d\u1eef li\u1ec7u", u"M\u00f4 t\u1ea3", u"R\u00e0ng bu\u1ed9c"]

    # --- users table ---
    elements.append(make_blank(doc))
    elements.append(make_normal(doc, u"B\u1ea3ng users - L\u01b0u th\u00f4ng tin t\u00e0i kho\u1ea3n ng\u01b0\u1eddi d\u00f9ng:"))
    elements.append(make_table(dd_headers, [
        ["id", "BIGINT UNSIGNED", u"M\u00e3 t\u00e0i kho\u1ea3n", "PRIMARY KEY, AUTO_INCREMENT"],
        ["email", "VARCHAR(255)", u"\u0110\u1ecba ch\u1ec9 email \u0111\u0103ng nh\u1eadp", "NOT NULL, UNIQUE"],
        ["password", "VARCHAR(255)", u"M\u1eadt kh\u1ea9u \u0111\u00e3 m\u00e3 h\u00f3a", "NOT NULL"],
        ["role", "TINYINT UNSIGNED", u"Vai tr\u00f2 (0: \u1ee9ng vi\u00ean, 1: NTD)", "NOT NULL"],
        ["is_active", "TINYINT(1)", u"Tr\u1ea1ng th\u00e1i ho\u1ea1t \u0111\u1ed9ng", "NOT NULL"],
        ["created_at", "TIMESTAMP", u"Th\u1eddi gian t\u1ea1o", "NULLABLE"],
        ["updated_at", "TIMESTAMP", u"Th\u1eddi gian c\u1eadp nh\u1eadt", "NULLABLE"],
    ]))
    elements.append(make_bang(doc, u"B\u1ea3ng 3.14a: C\u1ea5u tr\u00fac b\u1ea3ng users"))

    # --- candidates table ---
    elements.append(make_blank(doc))
    elements.append(make_normal(doc, u"B\u1ea3ng candidates - L\u01b0u th\u00f4ng tin \u1ee9ng vi\u00ean:"))
    elements.append(make_table(dd_headers, [
        ["id", "BIGINT UNSIGNED", u"M\u00e3 \u1ee9ng vi\u00ean", "PRIMARY KEY, AUTO_INCREMENT"],
        ["user_id", "BIGINT UNSIGNED", u"M\u00e3 t\u00e0i kho\u1ea3n li\u00ean k\u1ebft", u"FOREIGN KEY \u2192 users(id)"],
        ["firstname", "VARCHAR(50)", u"T\u00ean", "NOT NULL"],
        ["lastname", "VARCHAR(50)", u"H\u1ecd", "NOT NULL"],
        ["gender", "TINYINT UNSIGNED", u"Gi\u1edbi t\u00ednh (0: nam, 1: n\u1eef)", "NULLABLE"],
        ["dob", "DATE", u"Ng\u00e0y sinh", "NULLABLE"],
        ["phone", "CHAR(10)", u"S\u1ed1 \u0111i\u1ec7n tho\u1ea1i", "NULLABLE"],
        ["email", "VARCHAR(255)", u"Email li\u00ean h\u1ec7", "NOT NULL"],
        ["address", "VARCHAR(255)", u"\u0110\u1ecba ch\u1ec9", "NULLABLE"],
        ["link", "TEXT", u"Li\u00ean k\u1ebft c\u00e1 nh\u00e2n", "NULLABLE"],
        ["objective", "TEXT", u"M\u1ee5c ti\u00eau ngh\u1ec1 nghi\u1ec7p", "NULLABLE"],
        ["avatar", "TEXT", u"\u0110\u01b0\u1eddng d\u1eabn \u1ea3nh \u0111\u1ea1i di\u1ec7n", "NULLABLE"],
    ]))
    elements.append(make_bang(doc, u"B\u1ea3ng 3.14b: C\u1ea5u tr\u00fac b\u1ea3ng candidates"))

    # --- employers table ---
    elements.append(make_blank(doc))
    elements.append(make_normal(doc, u"B\u1ea3ng employers - L\u01b0u th\u00f4ng tin nh\u00e0 tuy\u1ec3n d\u1ee5ng:"))
    elements.append(make_table(dd_headers, [
        ["id", "BIGINT UNSIGNED", u"M\u00e3 nh\u00e0 tuy\u1ec3n d\u1ee5ng", "PRIMARY KEY, AUTO_INCREMENT"],
        ["user_id", "BIGINT UNSIGNED", u"M\u00e3 t\u00e0i kho\u1ea3n li\u00ean k\u1ebft", u"FOREIGN KEY \u2192 users(id)"],
        ["name", "VARCHAR(100)", u"T\u00ean c\u00f4ng ty", "NOT NULL"],
        ["address", "VARCHAR(255)", u"\u0110\u1ecba ch\u1ec9 tr\u1ee5 s\u1edf", "NOT NULL"],
        ["min_employees", "INT UNSIGNED", u"Quy m\u00f4 nh\u00e2n s\u1ef1 t\u1ed1i thi\u1ec3u", "NULLABLE"],
        ["max_employees", "INT UNSIGNED", u"Quy m\u00f4 nh\u00e2n s\u1ef1 t\u1ed1i \u0111a", "NULLABLE"],
        ["contact_name", "VARCHAR(60)", u"T\u00ean ng\u01b0\u1eddi li\u00ean h\u1ec7", "NULLABLE"],
        ["phone", "VARCHAR(15)", u"S\u1ed1 \u0111i\u1ec7n tho\u1ea1i", "NULLABLE"],
        ["website", "VARCHAR(255)", u"Website c\u00f4ng ty", "NULLABLE"],
        ["description", "LONGTEXT", u"M\u00f4 t\u1ea3 c\u00f4ng ty", "NULLABLE"],
        ["logo", "TEXT", u"\u0110\u01b0\u1eddng d\u1eabn logo", "NOT NULL"],
        ["image", "TEXT", u"\u0110\u01b0\u1eddng d\u1eabn \u1ea3nh b\u00eca", "NULLABLE"],
        ["is_hot", "TINYINT(1)", u"Nh\u00e0 tuy\u1ec3n d\u1ee5ng n\u1ed5i b\u1eadt", "NOT NULL"],
        ["is_active", "TINYINT(1)", u"Tr\u1ea1ng th\u00e1i ho\u1ea1t \u0111\u1ed9ng", "NOT NULL"],
    ]))
    elements.append(make_bang(doc, u"B\u1ea3ng 3.14c: C\u1ea5u tr\u00fac b\u1ea3ng employers"))

    # --- jobs table ---
    elements.append(make_blank(doc))
    elements.append(make_normal(doc, u"B\u1ea3ng jobs - L\u01b0u th\u00f4ng tin tin tuy\u1ec3n d\u1ee5ng:"))
    elements.append(make_table(dd_headers, [
        ["id", "BIGINT UNSIGNED", u"M\u00e3 tin tuy\u1ec3n d\u1ee5ng", "PRIMARY KEY, AUTO_INCREMENT"],
        ["employer_id", "BIGINT UNSIGNED", u"M\u00e3 nh\u00e0 tuy\u1ec3n d\u1ee5ng", u"FOREIGN KEY \u2192 employers(id)"],
        ["jtype_id", "BIGINT UNSIGNED", u"M\u00e3 lo\u1ea1i h\u00ecnh c\u00f4ng vi\u1ec7c", u"FOREIGN KEY \u2192 jtypes(id)"],
        ["jlevel_id", "BIGINT UNSIGNED", u"M\u00e3 c\u1ea5p b\u1eadc", u"FOREIGN KEY \u2192 jlevels(id)"],
        ["jname", "VARCHAR(150)", u"T\u00ean v\u1ecb tr\u00ed tuy\u1ec3n d\u1ee5ng", "NOT NULL"],
        ["address", "TEXT", u"\u0110\u1ecba ch\u1ec9 l\u00e0m vi\u1ec7c", "NOT NULL"],
        ["amount", "INT UNSIGNED", u"S\u1ed1 l\u01b0\u1ee3ng tuy\u1ec3n", "NULLABLE"],
        ["min_salary", "INT UNSIGNED", u"M\u1ee9c l\u01b0\u01a1ng t\u1ed1i thi\u1ec3u", "NULLABLE"],
        ["max_salary", "INT UNSIGNED", u"M\u1ee9c l\u01b0\u01a1ng t\u1ed1i \u0111a", "NULLABLE"],
        ["yoe", "TINYINT UNSIGNED", u"S\u1ed1 n\u0103m kinh nghi\u1ec7m y\u00eau c\u1ea7u", "NULLABLE"],
        ["gender", "TINYINT UNSIGNED", u"Gi\u1edbi t\u00ednh y\u00eau c\u1ea7u", "NULLABLE"],
        ["description", "LONGTEXT", u"M\u00f4 t\u1ea3 c\u00f4ng vi\u1ec7c", "NOT NULL"],
        ["expire_at", "DATE", u"H\u1ea1n n\u1ed9p h\u1ed3 s\u01a1", "NOT NULL"],
        ["is_hot", "TINYINT(1)", u"Tin n\u1ed5i b\u1eadt", "NOT NULL, DEFAULT 0"],
        ["is_active", "TINYINT(1)", u"Tr\u1ea1ng th\u00e1i tin", "NOT NULL, DEFAULT 1"],
    ]))
    elements.append(make_bang(doc, u"B\u1ea3ng 3.14d: C\u1ea5u tr\u00fac b\u1ea3ng jobs"))

    # --- resumes table ---
    elements.append(make_blank(doc))
    elements.append(make_normal(doc, u"B\u1ea3ng resumes - L\u01b0u th\u00f4ng tin CV/Resume tr\u1ef1c tuy\u1ebfn:"))
    elements.append(make_table(dd_headers, [
        ["id", "BIGINT UNSIGNED", u"M\u00e3 h\u1ed3 s\u01a1 CV", "PRIMARY KEY, AUTO_INCREMENT"],
        ["candidate_id", "BIGINT UNSIGNED", u"M\u00e3 \u1ee9ng vi\u00ean", u"FOREIGN KEY \u2192 candidates(id)"],
        ["title", "VARCHAR(60)", u"Ti\u00eau \u0111\u1ec1 CV", "NULLABLE"],
        ["fullname", "VARCHAR(100)", u"H\u1ecd t\u00ean tr\u00ean CV", "NULLABLE"],
        ["gender", "TINYINT UNSIGNED", u"Gi\u1edbi t\u00ednh", "NULLABLE"],
        ["dob", "DATE", u"Ng\u00e0y sinh", "NULLABLE"],
        ["phone", "CHAR(10)", u"S\u1ed1 \u0111i\u1ec7n tho\u1ea1i", "NULLABLE"],
        ["email", "VARCHAR(255)", "Email", "NULLABLE"],
        ["address", "VARCHAR(255)", u"\u0110\u1ecba ch\u1ec9", "NULLABLE"],
        ["objective", "TEXT", u"M\u1ee5c ti\u00eau ngh\u1ec1 nghi\u1ec7p", "NULLABLE"],
        ["avatar", "TEXT", u"\u1ea2nh \u0111\u1ea1i di\u1ec7n CV", "NULLABLE"],
        ["cv_link", "TEXT", u"Li\u00ean k\u1ebft CV \u0111\u00e3 xu\u1ea5t", "NULLABLE"],
        ["parts_order", "JSON", u"Th\u1ee9 t\u1ef1 c\u00e1c ph\u1ea7n trong CV", "NULLABLE"],
    ]))
    elements.append(make_bang(doc, u"B\u1ea3ng 3.14e: C\u1ea5u tr\u00fac b\u1ea3ng resumes"))

    # --- job_applying table ---
    elements.append(make_blank(doc))
    elements.append(make_normal(doc, u"B\u1ea3ng job_applying - L\u01b0u th\u00f4ng tin \u1ee9ng tuy\u1ec3n:"))
    elements.append(make_table(dd_headers, [
        ["job_id", "BIGINT UNSIGNED", u"M\u00e3 tin tuy\u1ec3n d\u1ee5ng", u"FOREIGN KEY \u2192 jobs(id)"],
        ["candidate_id", "BIGINT UNSIGNED", u"M\u00e3 \u1ee9ng vi\u00ean", u"FOREIGN KEY \u2192 candidates(id)"],
        ["cv_link", "TEXT", u"\u0110\u01b0\u1eddng d\u1eabn CV \u0111\u00e3 n\u1ed9p", "NULLABLE"],
        ["status", "TINYINT UNSIGNED", u"Tr\u1ea1ng th\u00e1i (0: ch\u1edd, 1: \u0111\u00e3 xem, 2: ch\u1ea5p nh\u1eadn, 3: t\u1eeb ch\u1ed1i)", "NOT NULL, DEFAULT 0"],
        ["created_at", "TIMESTAMP", u"Th\u1eddi gian \u1ee9ng tuy\u1ec3n", "NULLABLE"],
        ["updated_at", "TIMESTAMP", u"Th\u1eddi gian c\u1eadp nh\u1eadt", "NULLABLE"],
    ]))
    elements.append(make_bang(doc, u"B\u1ea3ng 3.14f: C\u1ea5u tr\u00fac b\u1ea3ng job_applying"))

    # --- saved_jobs table ---
    elements.append(make_blank(doc))
    elements.append(make_normal(doc, u"B\u1ea3ng saved_jobs - L\u01b0u tin tuy\u1ec3n d\u1ee5ng \u0111\u00e3 l\u01b0u:"))
    elements.append(make_table(dd_headers, [
        ["candidate_id", "BIGINT UNSIGNED", u"M\u00e3 \u1ee9ng vi\u00ean", u"FOREIGN KEY \u2192 candidates(id)"],
        ["job_id", "BIGINT UNSIGNED", u"M\u00e3 tin tuy\u1ec3n d\u1ee5ng", u"FOREIGN KEY \u2192 jobs(id)"],
    ]))
    elements.append(make_bang(doc, u"B\u1ea3ng 3.14g: C\u1ea5u tr\u00fac b\u1ea3ng saved_jobs"))

    # Insert all elements using addnext() from RM figure paragraph
    current = rm_para
    for elem in elements:
        current.addnext(elem)
        current = elem

    print(f"Inserted {len(elements)} elements total.")

    doc.save(DOC_PATH)
    print("Document saved successfully.")

    # === Verification ===
    print("\n=== Verification ===")
    doc2 = Document(DOC_PATH)

    # Check section headings exist
    found_324 = False
    found_325 = False
    idx_324 = -1
    idx_325 = -1
    idx_33 = -1
    for i, para in enumerate(doc2.paragraphs):
        if "3.2.4" in para.text and "Class Diagram" in para.text:
            found_324 = True
            idx_324 = i
        if "3.2.5" in para.text and "Data Dictionary" in para.text:
            found_325 = True
            idx_325 = i
        if para.style.name == 'Heading 2' and u'3.3' in para.text and u'ki\u1ec3m so\u00e1t' in para.text:
            idx_33 = i

    print(f"[{'OK' if found_324 else 'FAIL'}] Section 3.2.4 heading found at index {idx_324}")
    print(f"[{'OK' if found_325 else 'FAIL'}] Section 3.2.5 heading found at index {idx_325}")
    print(f"[{'OK' if idx_324 < idx_325 < idx_33 else 'FAIL'}] Order: 3.2.4({idx_324}) < 3.2.5({idx_325}) < 3.3({idx_33})")

    # Count tables
    table_count = len(doc2.tables)
    print(f"[INFO] Total tables in document: {table_count}")

    # Check for caption texts
    captions_found = 0
    expected_captions = [
        u"B\u1ea3ng 3.13",
        u"B\u1ea3ng 3.14a",
        u"B\u1ea3ng 3.14b",
        u"B\u1ea3ng 3.14c",
        u"B\u1ea3ng 3.14d",
        u"B\u1ea3ng 3.14e",
        u"B\u1ea3ng 3.14f",
        u"B\u1ea3ng 3.14g",
    ]
    for cap in expected_captions:
        found = False
        for p in doc2.paragraphs:
            if cap in p.text:
                found = True
                break
        if found:
            captions_found += 1
        else:
            print(f"  [FAIL] Caption '{cap}' not found")
    print(f"[{'OK' if captions_found == len(expected_captions) else 'FAIL'}] Captions found: {captions_found}/{len(expected_captions)}")


if __name__ == "__main__":
    main()
