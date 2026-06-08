import copy
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

DOC_PATH = 'D:/JobPotal/recruitment_web-xw1rpa/recruitment_web/K62_211213818_BuiDucTaiAnh_CNTT4.docx'
doc = Document(DOC_PATH)

# Find insertion point
insert_after = None
for i, p in enumerate(doc.paragraphs):
    if 'Bảng 2.12' in p.text and p.style.name == 'Bảng':
        insert_after = p._element
        print(f'Found insertion point at para {i}: {p.text}')
        break

if not insert_after:
    print('ERROR: Could not find Bảng 2.12')
    exit(1)

# Find reference table
ref_table = None
for t in doc.tables:
    if t.cell(0,0).text == 'Tên chức năng' and t.cell(0,1).text == 'Đăng nhập':
        ref_table = t
        break

def add_p_after(ref, text, style_name):
    new_p = OxmlElement('w:p')
    pPr = OxmlElement('w:pPr')
    pStyle = OxmlElement('w:pStyle')
    pStyle.set(qn('w:val'), style_name)
    pPr.append(pStyle)
    new_p.append(pPr)
    run = OxmlElement('w:r')
    t = OxmlElement('w:t')
    t.text = text
    t.set(qn('xml:space'), 'preserve')
    run.append(t)
    new_p.append(run)
    ref.addnext(new_p)
    return new_p

def add_spec_table_after(ref, data):
    new_tbl = copy.deepcopy(ref_table._tbl)
    rows = new_tbl.findall(qn('w:tr'))
    for r_idx, row in enumerate(rows):
        cells = row.findall(qn('w:tc'))
        if r_idx < len(data):
            for c_idx, cell in enumerate(cells):
                for p in cell.findall(qn('w:p')):
                    for run_el in p.findall(qn('w:r')):
                        for t_el in run_el.findall(qn('w:t')):
                            t_el.text = data[r_idx][c_idx]
    ref.addnext(new_tbl)
    return new_tbl

specs = [
    ('2.2.3.1. Quản lý học vấn', 'Bảng 2.13: Thông tin chung chức năng "Quản lý học vấn"', [
        ('Tên chức năng','Quản lý học vấn'),('Tác nhân','Ứng viên'),
        ('Mô tả','Cho phép ứng viên thêm, sửa, xóa thông tin học vấn bao gồm tên trường, chuyên ngành, thời gian học và mô tả chi tiết.'),
        ('Đầu vào','Tên trường, chuyên ngành, ngày bắt đầu, ngày kết thúc, mô tả'),
        ('Đầu ra','Thông tin học vấn được lưu/cập nhật/xóa thành công'),
        ('Điều kiện trước','Ứng viên đã đăng nhập vào hệ thống'),
        ('Điều kiện sau','Thông tin học vấn được cập nhật trong hồ sơ ứng viên'),
        ('Ngoại lệ','Thiếu thông tin bắt buộc'),('Các yêu cầu đặc biệt',''),
    ]),
    ('2.2.3.2. Quản lý kinh nghiệm làm việc', 'Bảng 2.14: Thông tin chung chức năng "Quản lý kinh nghiệm làm việc"', [
        ('Tên chức năng','Quản lý kinh nghiệm làm việc'),('Tác nhân','Ứng viên'),
        ('Mô tả','Cho phép ứng viên thêm, sửa, xóa kinh nghiệm làm việc bao gồm tên vị trí, công ty, loại hình công việc, thời gian và mô tả.'),
        ('Đầu vào','Tên vị trí, tên công ty, loại hình công việc, ngày bắt đầu, ngày kết thúc, mô tả'),
        ('Đầu ra','Thông tin kinh nghiệm được lưu/cập nhật/xóa thành công'),
        ('Điều kiện trước','Ứng viên đã đăng nhập vào hệ thống'),
        ('Điều kiện sau','Thông tin kinh nghiệm được cập nhật trong hồ sơ ứng viên'),
        ('Ngoại lệ','Thiếu thông tin bắt buộc'),('Các yêu cầu đặc biệt',''),
    ]),
    ('2.2.3.3. Quản lý kỹ năng', 'Bảng 2.15: Thông tin chung chức năng "Quản lý kỹ năng"', [
        ('Tên chức năng','Quản lý kỹ năng'),('Tác nhân','Ứng viên'),
        ('Mô tả','Cho phép ứng viên thêm, sửa, xóa các kỹ năng cá nhân bao gồm tên kỹ năng và mô tả mức độ thành thạo.'),
        ('Đầu vào','Tên kỹ năng, mô tả'),
        ('Đầu ra','Thông tin kỹ năng được lưu/cập nhật/xóa thành công'),
        ('Điều kiện trước','Ứng viên đã đăng nhập vào hệ thống'),
        ('Điều kiện sau','Thông tin kỹ năng được cập nhật trong hồ sơ ứng viên'),
        ('Ngoại lệ','Thiếu thông tin bắt buộc'),('Các yêu cầu đặc biệt',''),
    ]),
    ('2.2.3.4. Quản lý dự án cá nhân', 'Bảng 2.16: Thông tin chung chức năng "Quản lý dự án cá nhân"', [
        ('Tên chức năng','Quản lý dự án cá nhân'),('Tác nhân','Ứng viên'),
        ('Mô tả','Cho phép ứng viên thêm, sửa, xóa các dự án đã thực hiện bao gồm tên dự án, thời gian, mô tả, hình ảnh và liên kết.'),
        ('Đầu vào','Tên dự án, ngày bắt đầu, ngày kết thúc, mô tả, hình ảnh, liên kết'),
        ('Đầu ra','Thông tin dự án được lưu/cập nhật/xóa thành công'),
        ('Điều kiện trước','Ứng viên đã đăng nhập vào hệ thống'),
        ('Điều kiện sau','Thông tin dự án được cập nhật trong hồ sơ ứng viên'),
        ('Ngoại lệ','Thiếu thông tin bắt buộc'),('Các yêu cầu đặc biệt',''),
    ]),
    ('2.2.3.5. Quản lý chứng chỉ', 'Bảng 2.17: Thông tin chung chức năng "Quản lý chứng chỉ"', [
        ('Tên chức năng','Quản lý chứng chỉ'),('Tác nhân','Ứng viên'),
        ('Mô tả','Cho phép ứng viên thêm, sửa, xóa các chứng chỉ chuyên môn bao gồm tên chứng chỉ và hình ảnh minh chứng.'),
        ('Đầu vào','Tên chứng chỉ, hình ảnh'),
        ('Đầu ra','Thông tin chứng chỉ được lưu/cập nhật/xóa thành công'),
        ('Điều kiện trước','Ứng viên đã đăng nhập vào hệ thống'),
        ('Điều kiện sau','Thông tin chứng chỉ được cập nhật trong hồ sơ ứng viên'),
        ('Ngoại lệ','Thiếu thông tin bắt buộc'),('Các yêu cầu đặc biệt',''),
    ]),
    ('2.2.3.6. Quản lý giải thưởng', 'Bảng 2.18: Thông tin chung chức năng "Quản lý giải thưởng"', [
        ('Tên chức năng','Quản lý giải thưởng'),('Tác nhân','Ứng viên'),
        ('Mô tả','Cho phép ứng viên thêm, sửa, xóa các giải thưởng đã đạt được bao gồm tên giải thưởng và hình ảnh minh chứng.'),
        ('Đầu vào','Tên giải thưởng, hình ảnh'),
        ('Đầu ra','Thông tin giải thưởng được lưu/cập nhật/xóa thành công'),
        ('Điều kiện trước','Ứng viên đã đăng nhập vào hệ thống'),
        ('Điều kiện sau','Thông tin giải thưởng được cập nhật trong hồ sơ ứng viên'),
        ('Ngoại lệ','Thiếu thông tin bắt buộc'),('Các yêu cầu đặc biệt',''),
    ]),
    ('2.2.3.7. Quản lý hoạt động ngoại khóa', 'Bảng 2.19: Thông tin chung chức năng "Quản lý hoạt động ngoại khóa"', [
        ('Tên chức năng','Quản lý hoạt động ngoại khóa'),('Tác nhân','Ứng viên'),
        ('Mô tả','Cho phép ứng viên thêm, sửa, xóa các hoạt động ngoại khóa, tình nguyện bao gồm tên tổ chức, vai trò, thời gian, mô tả, hình ảnh và liên kết.'),
        ('Đầu vào','Tên tổ chức, vai trò, đang tham gia, ngày bắt đầu, ngày kết thúc, mô tả, hình ảnh, liên kết'),
        ('Đầu ra','Thông tin hoạt động được lưu/cập nhật/xóa thành công'),
        ('Điều kiện trước','Ứng viên đã đăng nhập vào hệ thống'),
        ('Điều kiện sau','Thông tin hoạt động được cập nhật trong hồ sơ ứng viên'),
        ('Ngoại lệ','Thiếu thông tin bắt buộc'),('Các yêu cầu đặc biệt',''),
    ]),
    ('2.2.3.8. Quản lý CV/Resume trực tuyến', 'Bảng 2.20: Thông tin chung chức năng "Quản lý CV/Resume trực tuyến"', [
        ('Tên chức năng','Quản lý CV/Resume trực tuyến'),('Tác nhân','Ứng viên'),
        ('Mô tả','Cho phép ứng viên tạo mới, chỉnh sửa, xóa CV trực tuyến. CV tổng hợp thông tin cá nhân, mục tiêu, học vấn, kinh nghiệm, kỹ năng, dự án, chứng chỉ, giải thưởng, hoạt động. Hỗ trợ tùy chỉnh thứ tự và xuất CV dạng hình ảnh.'),
        ('Đầu vào','Tiêu đề CV, thông tin cá nhân, mục tiêu nghề nghiệp, các phần nội dung CV'),
        ('Đầu ra','CV được tạo/cập nhật/xóa thành công'),
        ('Điều kiện trước','Ứng viên đã đăng nhập và có thông tin cá nhân cơ bản'),
        ('Điều kiện sau','CV được lưu trong hệ thống, sẵn sàng để sử dụng khi ứng tuyển'),
        ('Ngoại lệ','Thiếu thông tin bắt buộc'),
        ('Các yêu cầu đặc biệt','Hỗ trợ tùy chỉnh thứ tự các phần trong CV (parts_order)'),
    ]),
]

# Insert group heading
cur = insert_after
cur = add_p_after(cur, '2.2.3. Chức năng quản lý hồ sơ ứng viên', 'Heading 3')

for heading, caption, data in specs:
    cur = add_p_after(cur, heading, 'Heading 4')
    cur = add_p_after(cur, 'Thông tin chung chức năng', 'Normal')
    tbl = add_spec_table_after(cur, data)
    cur = tbl
    cur = add_p_after(cur, caption, 'Bảng')

doc.save(DOC_PATH)
print(f'Task 2 re-applied: added 2.2.3 with 8 sub-sections')
print(f'Total tables now: {len(Document(DOC_PATH).tables)}')
