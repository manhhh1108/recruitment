"""
Task 5: Fix figure numbering (Hình 2.6 -> Hình 3.1 in Chapter 3) and
add missing database entities to sections 3.2.1.1 and 3.2.1.2.
"""

from copy import deepcopy
from docx import Document
from lxml import etree

DOC_PATH = "D:/JobPotal/recruitment_web-xw1rpa/recruitment_web/K62_211213818_BuiDucTaiAnh_CNTT4.docx"

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def find_template_paragraph(doc, style_name):
    """Find a template paragraph with the given style to copy formatting from."""
    for para in doc.paragraphs:
        if para.style.name == style_name and para.runs:
            return para
    return None


def make_paragraph(doc, text):
    """Create a Normal paragraph by deep-copying an existing Normal paragraph's formatting."""
    template_para = find_template_paragraph(doc, "Normal")

    if template_para is not None:
        p = deepcopy(template_para._element)
        # Remove all runs and non-pPr children
        for child in list(p):
            tag = child.tag.split('}')[-1]
            if tag != 'pPr':
                p.remove(child)
        # Add new run with formatting from template
        r = etree.SubElement(p, f'{{{W}}}r')
        if template_para.runs:
            rPr_src = template_para.runs[0]._element.find(f'{{{W}}}rPr')
            if rPr_src is not None:
                new_rPr = deepcopy(rPr_src)
                r.insert(0, new_rPr)
        t = etree.SubElement(r, f'{{{W}}}t')
        t.text = text
        t.set(f'{{{W}}}space', "preserve")
    else:
        p = etree.Element(f'{{{W}}}p')
        pPr = etree.SubElement(p, f'{{{W}}}pPr')
        pStyle = etree.SubElement(pPr, f'{{{W}}}pStyle')
        pStyle.set(f'{{{W}}}val', "Normal")
        r = etree.SubElement(p, f'{{{W}}}r')
        t = etree.SubElement(r, f'{{{W}}}t')
        t.text = text
        t.set(f'{{{W}}}space', "preserve")

    return p


def fix_figure_numbering(doc):
    """Change 'Hình 2.6' to 'Hình 3.1' in paragraph [525] (Chapter 3 architecture diagram).

    The runs are split: 'Hình 2.' and '6' are separate runs.
    Only change the paragraph in Chapter 3 (about kiến trúc / Client - Server),
    NOT the one in Chapter 2 (about ca sử dụng).
    """
    changed = 0
    for para in doc.paragraphs:
        if 'Hình 2.6' not in para.text:
            continue
        # Skip the Use Case diagram paragraph (Chapter 2)
        if 'ca sử dụng' in para.text:
            continue
        # Skip TOC entries
        if para.style.name.startswith('toc'):
            continue

        # This should be the Chapter 3 architecture diagram paragraph
        # Handle split runs: combine adjacent runs that form "Hình 2.6"
        runs = para.runs
        i = 0
        while i < len(runs):
            run_text = runs[i].text
            if 'Hình 2.6' in run_text:
                # Simple case: full text in one run
                runs[i].text = run_text.replace('Hình 2.6', 'Hình 3.1')
                changed += 1
                break
            elif run_text.endswith('Hình 2.') and i + 1 < len(runs) and runs[i + 1].text.startswith('6'):
                # Split case: "Hình 2." + "6"
                runs[i].text = run_text[:-len('Hình 2.')] + 'Hình 3.'
                runs[i + 1].text = '1' + runs[i + 1].text[1:]
                changed += 1
                break
            i += 1

    print(f"Figure numbering: changed {changed} occurrence(s)")
    return changed


def add_entity_descriptions(doc):
    """Add missing entity descriptions after 'Tin tuyển dụng:' paragraph in section 3.2.1.1."""

    # Find the paragraph starting with "Tin tuyển dụng:" (index 539)
    insert_after = None
    for para in doc.paragraphs:
        if para.text.strip().startswith('Tin tuyển dụng:'):
            insert_after = para._element
            break

    if insert_after is None:
        raise RuntimeError("Could not find 'Tin tuyển dụng:' paragraph in section 3.2.1.1")

    intro_text = "Ngoài các thực thể đã nêu ở trên, hệ thống còn có các thực thể bổ sung sau đây phục vụ việc quản lý hồ sơ ứng viên, danh mục hệ thống và thông báo:"

    entity_texts = [
        "Học vấn: mã ứng viên, tên trường, chuyên ngành, ngày bắt đầu, ngày kết thúc, mô tả.",
        "Kinh nghiệm làm việc: mã ứng viên, mã loại công việc, tên vị trí, tên công ty, ngày bắt đầu, ngày kết thúc, mô tả.",
        "Kỹ năng ứng viên: mã ứng viên, tên kỹ năng, mô tả.",
        "Dự án cá nhân: mã ứng viên, tên dự án, ngày bắt đầu, ngày kết thúc, mô tả, hình ảnh, liên kết.",
        "Chứng chỉ: mã ứng viên, tên chứng chỉ, hình ảnh.",
        "Giải thưởng: mã ứng viên, tên giải thưởng, hình ảnh.",
        "Hoạt động ngoại khóa: mã ứng viên, tên tổ chức, vai trò, đang tham gia, ngày bắt đầu, ngày kết thúc, mô tả, hình ảnh, liên kết.",
        "Thông tin khác: mã ứng viên, tên mục, mô tả.",
        "Ngành nghề: tên ngành nghề.",
        "Cấp bậc: tên cấp bậc.",
        "Loại hình công việc: tên loại hình.",
        "Địa điểm: tên địa điểm.",
        "Hồ sơ CV: mã ứng viên, tiêu đề, họ tên, giới tính, ngày sinh, số điện thoại, email, địa chỉ, liên kết, ảnh đại diện, mục tiêu nghề nghiệp, tiêu đề các phần, liên kết CV, thứ tự các phần.",
        "Tin nhắn thông báo: mã ứng viên, mã tin tuyển dụng, nội dung, trạng thái đã đọc.",
    ]

    all_texts = [intro_text] + entity_texts

    current = insert_after
    for text in all_texts:
        p = make_paragraph(doc, text)
        current.addnext(p)
        current = p

    print(f"Entity descriptions: inserted {len(all_texts)} paragraphs after 'Tin tuyển dụng:'")


def add_key_identification(doc):
    """Add key identification entries after 'Tin tuyển dụng (Mã tin tuyển dụng...' paragraph in section 3.2.1.2."""

    # Find the paragraph: "Tin tuyển dụng (Mã tin tuyển dụng,...trạng thái)"
    insert_after = None
    for para in doc.paragraphs:
        if para.text.strip().startswith('Tin tuyển dụng (Mã tin') and 'trạng thái)' in para.text:
            insert_after = para._element
            break

    if insert_after is None:
        raise RuntimeError("Could not find 'Tin tuyển dụng (Mã tin tuyển dụng...' paragraph in section 3.2.1.2")

    key_texts = [
        'Xét Ngành nghề (tên ngành nghề). Ngành nghề chưa có thuộc tính định danh → Thêm "mã ngành nghề" → Khóa',
        "Ngành nghề (mã ngành nghề, tên ngành nghề)",
        'Xét Cấp bậc (tên cấp bậc). Cấp bậc chưa có thuộc tính định danh → Thêm "mã cấp bậc" → Khóa',
        "Cấp bậc (mã cấp bậc, tên cấp bậc)",
        'Xét Loại hình công việc (tên loại hình). Loại hình chưa có thuộc tính định danh → Thêm "mã loại hình" → Khóa',
        "Loại hình công việc (mã loại hình, tên loại hình)",
        'Xét Địa điểm (tên địa điểm). Địa điểm chưa có thuộc tính định danh → Thêm "mã địa điểm" → Khóa',
        "Địa điểm (mã địa điểm, tên địa điểm)",
        'Xét Hồ sơ CV (mã ứng viên, tiêu đề, họ tên, ...). Hồ sơ CV chưa có thuộc tính định danh → Thêm "mã hồ sơ CV" → Khóa',
        "Hồ sơ CV (mã hồ sơ CV, mã ứng viên, tiêu đề, họ tên, giới tính, ngày sinh, số điện thoại, email, địa chỉ, liên kết, ảnh đại diện, mục tiêu nghề nghiệp, tiêu đề các phần, liên kết CV, thứ tự các phần)",
        'Xét Tin nhắn thông báo (mã ứng viên, mã tin tuyển dụng, nội dung, trạng thái đã đọc). Tin nhắn chưa có thuộc tính định danh → Thêm "mã tin nhắn" → Khóa',
        "Tin nhắn thông báo (mã tin nhắn, mã ứng viên, mã tin tuyển dụng, nội dung, trạng thái đã đọc)",
    ]

    current = insert_after
    for text in key_texts:
        p = make_paragraph(doc, text)
        current.addnext(p)
        current = p

    print(f"Key identification: inserted {len(key_texts)} paragraphs after 'Tin tuyển dụng (Mã tin...)'")


def main():
    doc = Document(DOC_PATH)

    # Change 1: Fix figure numbering
    fix_figure_numbering(doc)

    # Change 2A: Add entity descriptions (section 3.2.1.1)
    add_entity_descriptions(doc)

    # Change 2B: Add key identification (section 3.2.1.2)
    add_key_identification(doc)

    doc.save(DOC_PATH)
    print("\nDocument saved successfully.")

    # Verification
    doc2 = Document(DOC_PATH)
    print("\n=== Verification ===")

    # Check figure numbering
    for para in doc2.paragraphs:
        if 'Hình 3.1' in para.text and para.style.name == 'hinh':
            print(f"[OK] Figure renamed: {para.text[:100]}")
        if 'Hình 2.6' in para.text:
            print(f"[INFO] Still has Hình 2.6: style={para.style.name} text={para.text[:100]}")

    # Check new entities in 3.2.1.1
    found_intro = False
    found_entities = 0
    for para in doc2.paragraphs:
        if 'thực thể bổ sung' in para.text:
            found_intro = True
            print(f"[OK] Intro paragraph found")
        if para.text.startswith('Học vấn:') or para.text.startswith('Tin nhắn thông báo:'):
            found_entities += 1
    print(f"[OK] Found {found_entities} boundary entity paragraphs (Học vấn + Tin nhắn thông báo)")

    # Check new entities in 3.2.1.2
    found_keys = 0
    for para in doc2.paragraphs:
        if para.text.startswith('Xét Ngành nghề') or para.text.startswith('Ngành nghề (mã'):
            found_keys += 1
        if para.text.startswith('Tin nhắn thông báo (mã tin nhắn'):
            found_keys += 1
    print(f"[OK] Found {found_keys} key identification entries (sample check)")


if __name__ == "__main__":
    main()
