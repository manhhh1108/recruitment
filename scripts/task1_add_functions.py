"""
Task 1: Add missing function items and descriptions to section 2.1.1
of K62_211213818_BuiDucTaiAnh_CNTT4.docx
"""

import copy
from docx import Document

DOC_PATH = "D:/JobPotal/recruitment_web-xw1rpa/recruitment_web/K62_211213818_BuiDucTaiAnh_CNTT4.docx"

# New list items to insert after paragraph 392
NEW_LIST_ITEMS = [
    "Quản lý học vấn",
    "Quản lý kinh nghiệm làm việc",
    "Quản lý kỹ năng",
    "Quản lý dự án cá nhân",
    "Quản lý chứng chỉ",
    "Quản lý giải thưởng",
    "Quản lý hoạt động ngoại khóa",
    "Quản lý thông tin khác",
    "Quản lý CV/Resume trực tuyến",
    "Xem thông báo kết quả ứng tuyển",
]

# New description paragraphs to insert after paragraph 409
NEW_DESCRIPTIONS = [
    "Chức năng quản lý học vấn: cho phép ứng viên thêm, sửa, xóa các thông tin về quá trình học tập của mình. Mỗi mục học vấn bao gồm tên trường, chuyên ngành, thời gian học và mô tả chi tiết. Thông tin học vấn được hiển thị trên hồ sơ xin việc và giúp nhà tuyển dụng đánh giá trình độ của ứng viên.",
    "Chức năng quản lý kinh nghiệm làm việc: cho phép ứng viên thêm, sửa, xóa các kinh nghiệm làm việc. Mỗi mục kinh nghiệm bao gồm tên vị trí, tên công ty, loại hình công việc, thời gian làm việc và mô tả công việc. Đây là thông tin quan trọng giúp nhà tuyển dụng đánh giá năng lực thực tế của ứng viên.",
    "Chức năng quản lý kỹ năng: cho phép ứng viên thêm, sửa, xóa các kỹ năng cá nhân. Mỗi mục kỹ năng bao gồm tên kỹ năng và mô tả mức độ thành thạo. Kỹ năng là tiêu chí quan trọng để nhà tuyển dụng lọc và đánh giá ứng viên phù hợp với yêu cầu công việc.",
    "Chức năng quản lý dự án cá nhân: cho phép ứng viên thêm, sửa, xóa các dự án đã tham gia hoặc thực hiện. Mỗi dự án bao gồm tên dự án, thời gian thực hiện, mô tả, hình ảnh minh họa và liên kết tham khảo. Thông tin dự án giúp ứng viên thể hiện năng lực thực tế ngoài bằng cấp và kinh nghiệm làm việc.",
    "Chức năng quản lý chứng chỉ: cho phép ứng viên thêm, sửa, xóa các chứng chỉ chuyên môn đã đạt được. Mỗi chứng chỉ bao gồm tên chứng chỉ và hình ảnh minh chứng. Chứng chỉ là bằng chứng khách quan về năng lực chuyên môn, giúp hồ sơ ứng viên thêm uy tín.",
    "Chức năng quản lý giải thưởng: cho phép ứng viên thêm, sửa, xóa các giải thưởng đã đạt được trong quá trình học tập và làm việc. Mỗi giải thưởng bao gồm tên giải thưởng và hình ảnh minh chứng. Giải thưởng giúp ứng viên nổi bật hơn trong mắt nhà tuyển dụng.",
    "Chức năng quản lý hoạt động ngoại khóa: cho phép ứng viên thêm, sửa, xóa các hoạt động ngoại khóa, tình nguyện hoặc cộng đồng. Mỗi hoạt động bao gồm tên tổ chức, vai trò, thời gian tham gia, mô tả, hình ảnh và liên kết. Hoạt động ngoại khóa giúp nhà tuyển dụng đánh giá kỹ năng mềm và tính cách của ứng viên.",
    "Chức năng quản lý thông tin khác: cho phép ứng viên bổ sung các thông tin cá nhân khác ngoài các mục đã có sẵn. Mỗi mục bao gồm tên mục và mô tả chi tiết. Chức năng này mang lại sự linh hoạt cho ứng viên trong việc trình bày hồ sơ.",
    "Chức năng quản lý CV/Resume trực tuyến: cho phép ứng viên tạo mới, chỉnh sửa và xóa các CV trực tuyến trên hệ thống. Mỗi CV bao gồm thông tin cá nhân, mục tiêu nghề nghiệp, học vấn, kinh nghiệm, kỹ năng, dự án, chứng chỉ, giải thưởng và hoạt động. Ứng viên có thể tạo nhiều CV khác nhau để phù hợp với từng vị trí ứng tuyển. Hệ thống hỗ trợ tùy chỉnh thứ tự các phần trong CV và xuất CV dưới dạng hình ảnh.",
    "Chức năng xem thông báo kết quả ứng tuyển: cho phép ứng viên nhận và xem các thông báo từ hệ thống về kết quả xử lý hồ sơ ứng tuyển. Khi nhà tuyển dụng thay đổi trạng thái hồ sơ (đã xem, chấp nhận, từ chối, mời phỏng vấn), hệ thống tự động gửi thông báo cho ứng viên. Chức năng này giúp ứng viên chủ động theo dõi tiến trình ứng tuyển của mình.",
]


def insert_paragraph_after(ref_paragraph_element, new_element):
    """Insert new_element after ref_paragraph_element in the XML tree."""
    ref_paragraph_element.addnext(new_element)


def clone_and_replace_text(source_paragraph, new_text):
    """
    Deep copy the XML element of source_paragraph, then replace all run text
    with new_text while preserving formatting.
    """
    from lxml import etree
    new_elem = copy.deepcopy(source_paragraph._element)

    # Find the namespace
    nsmap = source_paragraph._element.nsmap
    w_ns = nsmap.get('w', 'http://schemas.openxmlformats.org/wordprocessingml/2006/main')

    # Find all runs and their text elements
    runs = new_elem.findall(f'{{{w_ns}}}r')

    if runs:
        # Clear text from all runs except the first one
        for i, run in enumerate(runs):
            t_elements = run.findall(f'{{{w_ns}}}t')
            if i == 0 and t_elements:
                t_elements[0].text = new_text
                # Preserve spaces
                t_elements[0].set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
                # Remove extra t elements in first run
                for t in t_elements[1:]:
                    run.remove(t)
            else:
                # Remove all t elements from subsequent runs
                for t in t_elements:
                    run.remove(t)
                # Remove the run if it has no text (but keep if it has other properties)
                # Actually, remove entire extra runs to avoid empty content
                if not run.findall(f'{{{w_ns}}}t'):
                    new_elem.remove(run)
    return new_elem


def main():
    doc = Document(DOC_PATH)
    paragraphs = doc.paragraphs
    total_before = len(paragraphs)

    print(f"Total paragraphs before: {total_before}")

    # Verify paragraph 392 content
    p392 = paragraphs[392]
    print(f"Paragraph 392 text: '{p392.text}'")
    print(f"Paragraph 392 style: '{p392.style.name}'")
    assert "Tìm kiếm bài viết" in p392.text, f"Expected 'Tìm kiếm bài viết' at index 392, got '{p392.text}'"

    # Verify paragraph 409 content
    p409 = paragraphs[409]
    print(f"Paragraph 409 text: '{p409.text[:60]}...'")
    print(f"Paragraph 409 style: '{p409.style.name}'")
    assert "Chức năng tìm kiếm bài viết" in p409.text, f"Expected description at index 409, got '{p409.text[:60]}'"

    # --- Insert list items after paragraph 392 ---
    # Insert in reverse order so they end up in correct order
    ref_elem_392 = p392._element
    for item_text in reversed(NEW_LIST_ITEMS):
        new_elem = clone_and_replace_text(p392, item_text)
        insert_paragraph_after(ref_elem_392, new_elem)
    print(f"Inserted {len(NEW_LIST_ITEMS)} list items after paragraph 392")

    # --- Insert descriptions after paragraph 409 ---
    # After inserting 10 list items above, paragraph 409 shifted to 419
    # But we already have the reference to the element, so we use p409._element
    ref_elem_409 = p409._element
    for desc_text in reversed(NEW_DESCRIPTIONS):
        new_elem = clone_and_replace_text(p409, desc_text)
        insert_paragraph_after(ref_elem_409, new_elem)
    print(f"Inserted {len(NEW_DESCRIPTIONS)} description paragraphs after paragraph 409 (original index)")

    # Save
    doc.save(DOC_PATH)
    print(f"Document saved to {DOC_PATH}")

    # Verify
    doc2 = Document(DOC_PATH)
    total_after = len(doc2.paragraphs)
    print(f"\nTotal paragraphs after: {total_after}")
    print(f"Expected increase: {len(NEW_LIST_ITEMS) + len(NEW_DESCRIPTIONS)} = 20")
    print(f"Actual increase: {total_after - total_before}")

    # Verify list items
    print("\n--- Verifying list items (paragraphs 393-402) ---")
    for i in range(393, 403):
        p = doc2.paragraphs[i]
        print(f"  [{i}] style='{p.style.name}' text='{p.text}'")

    # Verify descriptions (original 409 + 10 list items = 419, so new descs at 420-429)
    print("\n--- Verifying descriptions (paragraphs 420-429) ---")
    for i in range(420, 430):
        p = doc2.paragraphs[i]
        print(f"  [{i}] style='{p.style.name}' text='{p.text[:70]}...'")


if __name__ == "__main__":
    main()
