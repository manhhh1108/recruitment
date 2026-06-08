"""
Task 2: Add function specification group 2.2.3 to the docx document.
Inserts heading + 8 sub-sections (each with heading, normal text, caption, and 9x2 table)
after the "Bảng 2.12" caption paragraph.
"""

from copy import deepcopy
from docx import Document
from lxml import etree

DOC_PATH = "D:/JobPotal/recruitment_web-xw1rpa/recruitment_web/K62_211213818_BuiDucTaiAnh_CNTT4.docx"

# Sub-section data: (heading_number, heading_text, table_caption, table_rows)
# Each table_rows is list of 9 (label, value) pairs
SUBSECTIONS = [
    (
        "2.2.3.1. Quản lý học vấn",
        'Bảng 2.13: Thông tin chung chức năng "Quản lý học vấn"',
        [
            ("Tên chức năng", "Quản lý học vấn"),
            ("Tác nhân", "Ứng viên"),
            ("Mô tả", "Cho phép ứng viên thêm, sửa, xóa thông tin học vấn bao gồm tên trường, chuyên ngành, thời gian học và mô tả chi tiết."),
            ("Đầu vào", "Tên trường, chuyên ngành, ngày bắt đầu, ngày kết thúc, mô tả"),
            ("Đầu ra", "Thông tin học vấn được lưu/cập nhật/xóa thành công"),
            ("Điều kiện trước", "Ứng viên đã đăng nhập vào hệ thống"),
            ("Điều kiện sau", "Thông tin học vấn được cập nhật trong hồ sơ ứng viên"),
            ("Ngoại lệ", "Thiếu thông tin bắt buộc"),
            ("Các yêu cầu đặc biệt", ""),
        ],
    ),
    (
        "2.2.3.2. Quản lý kinh nghiệm làm việc",
        'Bảng 2.14: Thông tin chung chức năng "Quản lý kinh nghiệm làm việc"',
        [
            ("Tên chức năng", "Quản lý kinh nghiệm làm việc"),
            ("Tác nhân", "Ứng viên"),
            ("Mô tả", "Cho phép ứng viên thêm, sửa, xóa kinh nghiệm làm việc bao gồm tên vị trí, công ty, loại hình công việc, thời gian và mô tả."),
            ("Đầu vào", "Tên vị trí, tên công ty, loại hình công việc, ngày bắt đầu, ngày kết thúc, mô tả"),
            ("Đầu ra", "Thông tin kinh nghiệm được lưu/cập nhật/xóa thành công"),
            ("Điều kiện trước", "Ứng viên đã đăng nhập vào hệ thống"),
            ("Điều kiện sau", "Thông tin kinh nghiệm được cập nhật trong hồ sơ ứng viên"),
            ("Ngoại lệ", "Thiếu thông tin bắt buộc"),
            ("Các yêu cầu đặc biệt", ""),
        ],
    ),
    (
        "2.2.3.3. Quản lý kỹ năng",
        'Bảng 2.15: Thông tin chung chức năng "Quản lý kỹ năng"',
        [
            ("Tên chức năng", "Quản lý kỹ năng"),
            ("Tác nhân", "Ứng viên"),
            ("Mô tả", "Cho phép ứng viên thêm, sửa, xóa các kỹ năng cá nhân bao gồm tên kỹ năng và mô tả mức độ thành thạo."),
            ("Đầu vào", "Tên kỹ năng, mô tả"),
            ("Đầu ra", "Thông tin kỹ năng được lưu/cập nhật/xóa thành công"),
            ("Điều kiện trước", "Ứng viên đã đăng nhập vào hệ thống"),
            ("Điều kiện sau", "Thông tin kỹ năng được cập nhật trong hồ sơ ứng viên"),
            ("Ngoại lệ", "Thiếu thông tin bắt buộc"),
            ("Các yêu cầu đặc biệt", ""),
        ],
    ),
    (
        "2.2.3.4. Quản lý dự án cá nhân",
        'Bảng 2.16: Thông tin chung chức năng "Quản lý dự án cá nhân"',
        [
            ("Tên chức năng", "Quản lý dự án cá nhân"),
            ("Tác nhân", "Ứng viên"),
            ("Mô tả", "Cho phép ứng viên thêm, sửa, xóa các dự án đã thực hiện bao gồm tên dự án, thời gian, mô tả, hình ảnh và liên kết."),
            ("Đầu vào", "Tên dự án, ngày bắt đầu, ngày kết thúc, mô tả, hình ảnh, liên kết"),
            ("Đầu ra", "Thông tin dự án được lưu/cập nhật/xóa thành công"),
            ("Điều kiện trước", "Ứng viên đã đăng nhập vào hệ thống"),
            ("Điều kiện sau", "Thông tin dự án được cập nhật trong hồ sơ ứng viên"),
            ("Ngoại lệ", "Thiếu thông tin bắt buộc"),
            ("Các yêu cầu đặc biệt", ""),
        ],
    ),
    (
        "2.2.3.5. Quản lý chứng chỉ",
        'Bảng 2.17: Thông tin chung chức năng "Quản lý chứng chỉ"',
        [
            ("Tên chức năng", "Quản lý chứng chỉ"),
            ("Tác nhân", "Ứng viên"),
            ("Mô tả", "Cho phép ứng viên thêm, sửa, xóa các chứng chỉ chuyên môn bao gồm tên chứng chỉ và hình ảnh minh chứng."),
            ("Đầu vào", "Tên chứng chỉ, hình ảnh"),
            ("Đầu ra", "Thông tin chứng chỉ được lưu/cập nhật/xóa thành công"),
            ("Điều kiện trước", "Ứng viên đã đăng nhập vào hệ thống"),
            ("Điều kiện sau", "Thông tin chứng chỉ được cập nhật trong hồ sơ ứng viên"),
            ("Ngoại lệ", "Thiếu thông tin bắt buộc"),
            ("Các yêu cầu đặc biệt", ""),
        ],
    ),
    (
        "2.2.3.6. Quản lý giải thưởng",
        'Bảng 2.18: Thông tin chung chức năng "Quản lý giải thưởng"',
        [
            ("Tên chức năng", "Quản lý giải thưởng"),
            ("Tác nhân", "Ứng viên"),
            ("Mô tả", "Cho phép ứng viên thêm, sửa, xóa các giải thưởng đã đạt được bao gồm tên giải thưởng và hình ảnh minh chứng."),
            ("Đầu vào", "Tên giải thưởng, hình ảnh"),
            ("Đầu ra", "Thông tin giải thưởng được lưu/cập nhật/xóa thành công"),
            ("Điều kiện trước", "Ứng viên đã đăng nhập vào hệ thống"),
            ("Điều kiện sau", "Thông tin giải thưởng được cập nhật trong hồ sơ ứng viên"),
            ("Ngoại lệ", "Thiếu thông tin bắt buộc"),
            ("Các yêu cầu đặc biệt", ""),
        ],
    ),
    (
        "2.2.3.7. Quản lý hoạt động ngoại khóa",
        'Bảng 2.19: Thông tin chung chức năng "Quản lý hoạt động ngoại khóa"',
        [
            ("Tên chức năng", "Quản lý hoạt động ngoại khóa"),
            ("Tác nhân", "Ứng viên"),
            ("Mô tả", "Cho phép ứng viên thêm, sửa, xóa các hoạt động ngoại khóa, tình nguyện bao gồm tên tổ chức, vai trò, thời gian, mô tả, hình ảnh và liên kết."),
            ("Đầu vào", "Tên tổ chức, vai trò, đang tham gia, ngày bắt đầu, ngày kết thúc, mô tả, hình ảnh, liên kết"),
            ("Đầu ra", "Thông tin hoạt động được lưu/cập nhật/xóa thành công"),
            ("Điều kiện trước", "Ứng viên đã đăng nhập vào hệ thống"),
            ("Điều kiện sau", "Thông tin hoạt động được cập nhật trong hồ sơ ứng viên"),
            ("Ngoại lệ", "Thiếu thông tin bắt buộc"),
            ("Các yêu cầu đặc biệt", ""),
        ],
    ),
    (
        "2.2.3.8. Quản lý CV/Resume trực tuyến",
        'Bảng 2.20: Thông tin chung chức năng "Quản lý CV/Resume trực tuyến"',
        [
            ("Tên chức năng", "Quản lý CV/Resume trực tuyến"),
            ("Tác nhân", "Ứng viên"),
            ("Mô tả", "Cho phép ứng viên tạo mới, chỉnh sửa, xóa CV trực tuyến. CV tổng hợp thông tin cá nhân, mục tiêu, học vấn, kinh nghiệm, kỹ năng, dự án, chứng chỉ, giải thưởng, hoạt động. Hỗ trợ tùy chỉnh thứ tự và xuất CV dạng hình ảnh."),
            ("Đầu vào", "Tiêu đề CV, thông tin cá nhân, mục tiêu nghề nghiệp, các phần nội dung CV"),
            ("Đầu ra", "CV được tạo/cập nhật/xóa thành công"),
            ("Điều kiện trước", "Ứng viên đã đăng nhập và có thông tin cá nhân cơ bản"),
            ("Điều kiện sau", "CV được lưu trong hệ thống, sẵn sàng để sử dụng khi ứng tuyển"),
            ("Ngoại lệ", "Thiếu thông tin bắt buộc"),
            ("Các yêu cầu đặc biệt", "Hỗ trợ tùy chỉnh thứ tự các phần trong CV (parts_order)"),
        ],
    ),
]

NSMAP = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}


def get_style_id(doc, style_name):
    """Get the style_id for a given style display name."""
    for s in doc.styles:
        if s.name == style_name:
            return s.style_id
    return style_name  # fallback


def find_template_paragraph(doc, style_name):
    """Find a template paragraph with the given style to copy formatting from."""
    for para in doc.paragraphs:
        if para.style.name == style_name and para.runs:
            return para
    return None


def make_paragraph(style_name, text, doc):
    """Create a new paragraph XML element with the given style and text."""
    W = NSMAP["w"]

    # Find a template paragraph to clone formatting from
    template_para = find_template_paragraph(doc, style_name)

    if template_para is not None:
        # Deep copy the entire template paragraph and replace text
        p = deepcopy(template_para._element)
        # Clear all runs
        for old_r in p.findall(f'{{{W}}}r'):
            p.remove(old_r)
        # Clear any bookmarks or other non-pPr children
        for child in list(p):
            tag = child.tag.split('}')[-1]
            if tag != 'pPr':
                p.remove(child)
        # Add new run with text
        r = etree.SubElement(p, f'{{{W}}}r')
        # Copy rPr from template
        if template_para.runs:
            rPr_src = template_para.runs[0]._element.find(f'{{{W}}}rPr')
            if rPr_src is not None:
                r.insert(0, deepcopy(rPr_src))
        t = etree.SubElement(r, f'{{{W}}}t')
        t.text = text
        t.set(f'{{{W}}}space', "preserve")
    else:
        # Fallback: create from scratch
        p = etree.Element(f'{{{W}}}p')
        pPr = etree.SubElement(p, f'{{{W}}}pPr')
        pStyle = etree.SubElement(pPr, f'{{{W}}}pStyle')
        pStyle.set(f'{{{W}}}val', get_style_id(doc, style_name))
        r = etree.SubElement(p, f'{{{W}}}r')
        t = etree.SubElement(r, f'{{{W}}}t')
        t.text = text
        t.set(f'{{{W}}}space', "preserve")

    return p


def make_table_from_template(template_table, rows_data):
    """Deep copy a template table and replace cell text while preserving formatting."""
    W = NSMAP["w"]
    tbl_xml = deepcopy(template_table._tbl)

    # Get all rows
    tbl_rows = tbl_xml.findall(f'{{{W}}}tr')
    assert len(tbl_rows) == 9, f"Expected 9 rows, got {len(tbl_rows)}"

    for row_idx, (label, value) in enumerate(rows_data):
        tr = tbl_rows[row_idx]
        tcs = tr.findall(f'{{{W}}}tc')
        assert len(tcs) == 2, f"Expected 2 cells, got {len(tcs)}"

        for cell_idx, cell_text in enumerate([label, value]):
            tc = tcs[cell_idx]
            paragraphs = tc.findall(f'{{{W}}}p')

            if paragraphs:
                # Keep the first paragraph, remove extras
                first_p = paragraphs[0]
                for extra_p in paragraphs[1:]:
                    tc.remove(extra_p)

                # Preserve pPr (paragraph properties) but replace runs
                for old_r in first_p.findall(f'{{{W}}}r'):
                    first_p.remove(old_r)

                # Get rPr from original first run if available
                rPr_template = None
                orig_tc = template_table._tbl.findall(f'{{{W}}}tr')[row_idx].findall(f'{{{W}}}tc')[cell_idx]
                orig_p = orig_tc.findall(f'{{{W}}}p')
                if orig_p:
                    orig_runs = orig_p[0].findall(f'{{{W}}}r')
                    if orig_runs:
                        rPr_template = orig_runs[0].find(f'{{{W}}}rPr')

                r = etree.SubElement(first_p, f'{{{W}}}r')
                if rPr_template is not None:
                    r.insert(0, deepcopy(rPr_template))
                t_elem = etree.SubElement(r, f'{{{W}}}t')
                t_elem.text = cell_text
                t_elem.set(f'{{{W}}}space', "preserve")
            else:
                # Fallback: create new paragraph
                p = etree.SubElement(tc, f'{{{W}}}p')
                r = etree.SubElement(p, f'{{{W}}}r')
                t_elem = etree.SubElement(r, f'{{{W}}}t')
                t_elem.text = cell_text
                t_elem.set(f'{{{W}}}space', "preserve")

    return tbl_xml


def main():
    doc = Document(DOC_PATH)

    # Find insertion point: paragraph containing "Bảng 2.12" with style "Bảng"
    body = doc.element.body
    insert_after = None
    for para in doc.paragraphs:
        if "Bảng 2.12" in para.text and para.style.name == "Bảng":
            insert_after = para._element
            break

    if insert_after is None:
        raise RuntimeError("Could not find 'Bảng 2.12' caption paragraph with style 'Bảng'")

    # Template table (table index 1 = 9x2 spec table)
    template_table = doc.tables[1]

    # Track the current element to insert after
    current = insert_after

    # First: add the section heading (Heading 3)
    h3 = make_paragraph("Heading 3", "2.2.3. Chức năng quản lý hồ sơ ứng viên", doc)
    current.addnext(h3)
    current = h3

    # Add each sub-section
    for heading_text, caption_text, table_rows in SUBSECTIONS:
        # Heading 4
        h4 = make_paragraph("Heading 4", heading_text, doc)
        current.addnext(h4)
        current = h4

        # Normal paragraph: "Thông tin chung chức năng"
        normal_p = make_paragraph("Normal", "Thông tin chung chức năng", doc)
        current.addnext(normal_p)
        current = normal_p

        # Caption paragraph (style "Bảng")
        caption_p = make_paragraph("Bảng", caption_text, doc)
        current.addnext(caption_p)
        current = caption_p

        # Table
        tbl = make_table_from_template(template_table, table_rows)
        current.addnext(tbl)
        current = tbl

    doc.save(DOC_PATH)
    print("Document saved successfully.")

    # Verify
    doc2 = Document(DOC_PATH)
    print(f"Total tables: {len(doc2.tables)}")
    # Check new tables exist
    for para in doc2.paragraphs:
        if "2.2.3" in para.text:
            print(f"  Found: [{para.style.name}] {para.text[:100]}")


if __name__ == "__main__":
    main()
