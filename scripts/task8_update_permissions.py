"""
Task 8: Update permission tables in section 3.3.
Add new rows to both the data permission table and function permission table.
"""

from copy import deepcopy
from docx import Document
from docx.oxml.ns import qn

DOC_PATH = "D:/JobPotal/recruitment_web-xw1rpa/recruitment_web/K62_211213818_BuiDucTaiAnh_CNTT4.docx"

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

# New rows for the data permission table
DATA_ROWS = [
    ["tbl_HocVan",       "C, R, E, D", "",           "C, R, E, D"],
    ["tbl_KinhNghiem",   "C, R, E, D", "",           "C, R, E, D"],
    ["tbl_KyNang",       "C, R, E, D", "",           "C, R, E, D"],
    ["tbl_DuAn",         "C, R, E, D", "",           "C, R, E, D"],
    ["tbl_ChungChi",     "C, R, E, D", "",           "C, R, E, D"],
    ["tbl_GiaiThuong",   "C, R, E, D", "",           "C, R, E, D"],
    ["tbl_HoatDong",     "C, R, E, D", "",           "C, R, E, D"],
    ["tbl_ThongTinKhac", "C, R, E, D", "",           "C, R, E, D"],
    ["tbl_HoSoCV",       "C, R, E, D", "",           "C, R, E, D"],
    ["tbl_NganhNghe",    "R",          "R",          "C, R, E, D"],
    ["tbl_CapBac",       "R",          "R",          "C, R, E, D"],
    ["tbl_LoaiHinh",     "R",          "R",          "C, R, E, D"],
    ["tbl_DiaDiem",      "R",          "R",          "C, R, E, D"],
    ["tbl_TinNhan",      "R",          "",           "C, R, E, D"],
    ["tbl_UngTuyen",     "C, R",       "R, E",       "C, R, E, D"],
    ["tbl_LuuTin",       "C, R, D",    "",           "C, R, E, D"],
]

# New rows for the function permission table
FUNC_ROWS = [
    ["Quản lý học vấn",            "A",     "not A", "A"],
    ["Quản lý kinh nghiệm",       "A",     "not A", "A"],
    ["Quản lý kỹ năng",           "A",     "not A", "A"],
    ["Quản lý dự án",             "A",     "not A", "A"],
    ["Quản lý chứng chỉ",         "A",     "not A", "A"],
    ["Quản lý giải thưởng",       "A",     "not A", "A"],
    ["Quản lý hoạt động",         "A",     "not A", "A"],
    ["Quản lý thông tin khác",    "A",     "not A", "A"],
    ["Quản lý CV/Resume",         "A",     "not A", "A"],
    ["Xem thông báo kết quả",     "A",     "not A", "A"],
    ["Tìm kiếm nhà tuyển dụng",  "A",     "A",     "A"],
]


def get_cell_text(cell):
    """Get the text content of a table cell."""
    return cell.text.strip()


def set_cell_text(cell, text):
    """Replace all text in a cell, preserving formatting from the first paragraph/run."""
    # Clear all paragraphs except the first
    paragraphs = cell._element.findall(qn('w:p'))
    for p in paragraphs[1:]:
        cell._element.remove(p)

    # In the first paragraph, clear all runs except the first
    first_p = paragraphs[0]
    runs = first_p.findall(qn('w:r'))
    for r in runs[1:]:
        first_p.remove(r)

    if runs:
        # Set text on the first run
        t_elem = runs[0].find(qn('w:t'))
        if t_elem is None:
            t_elem = runs[0].makeelement(qn('w:t'), {})
            runs[0].append(t_elem)
        t_elem.text = text
        t_elem.set(qn('xml:space'), 'preserve')
    else:
        # Create a run with the text
        r = first_p.makeelement(qn('w:r'), {})
        t_elem = r.makeelement(qn('w:t'), {})
        t_elem.text = text
        t_elem.set(qn('xml:space'), 'preserve')
        r.append(t_elem)
        first_p.append(r)


def find_tables(doc):
    """Find the data permission table and function permission table."""
    data_table = None
    func_table = None

    for i, table in enumerate(doc.tables):
        rows = table.rows
        if len(rows) < 3 or len(table.columns) != 4:
            continue

        # Check first cell of row 0
        first_cell = get_cell_text(rows[0].cells[0])
        if "Bảng dữ liệu" not in first_cell:
            continue

        # Check row 2 to distinguish the two tables
        row2_text = get_cell_text(rows[2].cells[0])
        if "tbl_TaiKhoan" in row2_text:
            data_table = table
            print(f"Found data permission table at index {i} with {len(rows)} rows")
        elif "Đăng nhập" in row2_text:
            func_table = table
            print(f"Found function permission table at index {i} with {len(rows)} rows")

    return data_table, func_table


def add_rows_to_table(table, new_rows):
    """Add new rows to a table by deep-copying the last row and replacing text."""
    tbl_elem = table._tbl
    last_row = table.rows[-1]._tr

    for row_data in new_rows:
        new_tr = deepcopy(last_row)
        # Get cells in the new row
        cells = new_tr.findall(qn('w:tc'))
        for j, cell_xml in enumerate(cells):
            # Build a temporary cell wrapper to use set_cell_text
            # We'll manipulate XML directly instead
            paragraphs = cell_xml.findall(qn('w:p'))
            # Clear all paragraphs except first
            for p in paragraphs[1:]:
                cell_xml.remove(p)

            first_p = paragraphs[0] if paragraphs else None
            if first_p is None:
                first_p = cell_xml.makeelement(qn('w:p'), {})
                cell_xml.append(first_p)

            # Clear all runs except first
            runs = first_p.findall(qn('w:r'))
            for r in runs[1:]:
                first_p.remove(r)

            text = row_data[j] if j < len(row_data) else ""

            if runs:
                t_elem = runs[0].find(qn('w:t'))
                if t_elem is None:
                    t_elem = runs[0].makeelement(qn('w:t'), {})
                    runs[0].append(t_elem)
                t_elem.text = text
                t_elem.set(qn('xml:space'), 'preserve')
            else:
                r = first_p.makeelement(qn('w:r'), {})
                t_elem = r.makeelement(qn('w:t'), {})
                t_elem.text = text
                t_elem.set(qn('xml:space'), 'preserve')
                r.append(t_elem)
                first_p.append(r)

        tbl_elem.append(new_tr)


def main():
    doc = Document(DOC_PATH)
    print(f"Total tables in document: {len(doc.tables)}")

    data_table, func_table = find_tables(doc)

    if data_table is None:
        print("ERROR: Could not find data permission table!")
        return
    if func_table is None:
        print("ERROR: Could not find function permission table!")
        return

    data_rows_before = len(data_table.rows)
    func_rows_before = len(func_table.rows)

    add_rows_to_table(data_table, DATA_ROWS)
    add_rows_to_table(func_table, FUNC_ROWS)

    # Re-read to verify
    doc.save(DOC_PATH)
    doc2 = Document(DOC_PATH)
    data_table2, func_table2 = find_tables(doc2)

    data_rows_after = len(data_table2.rows)
    func_rows_after = len(func_table2.rows)

    print(f"\nData permission table: {data_rows_before} -> {data_rows_after} rows (added {data_rows_after - data_rows_before})")
    print(f"Function permission table: {func_rows_before} -> {func_rows_after} rows (added {func_rows_after - func_rows_before})")

    # Verify content of new rows
    print("\n--- Data table new rows ---")
    for i in range(data_rows_before, data_rows_after):
        row = data_table2.rows[i]
        cells = [c.text.strip() for c in row.cells]
        print(f"  Row {i}: {cells}")

    print("\n--- Function table new rows ---")
    for i in range(func_rows_before, func_rows_after):
        row = func_table2.rows[i]
        cells = [c.text.strip() for c in row.cells]
        print(f"  Row {i}: {cells}")

    print("\nDone!")


if __name__ == "__main__":
    main()
